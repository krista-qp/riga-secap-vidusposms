#!/usr/bin/env python3
"""Rīgas SECAP vidusposma pārskata vietne – datu sagatavošana.

Nolasa divas Excel datnes mapē avoti/ un izveido:
  assets/data.js                 – visi vietnes dati (rādītāji, mērķi, pasākumi, teksti)
  dati/*.csv, dati/*.xlsx        – lejupielādējamie atvērtie dati

Lietošana (no repozitorija saknes):
  python3 build/build.py
  python3 build/build.py --raditaji ceļš/uz/Raditaji.xlsx --saturs ceļš/uz/saturs.xlsx

Nepieciešams: Python 3.9+, openpyxl (pip install openpyxl).
Rādītāju datnē tiek lasītas saglabātās (aprēķinātās) vērtības – pēc izmaiņām
Excel datni saglabājiet ar Excel, lai formulu vērtības būtu atjaunotas.
"""
import argparse, csv, datetime, json, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
YEARS = [2020, 2021, 2022, 2023, 2024, 2025]

# ---------------------------------------------------------------- konfigurācija
# Mērķu sasaiste ar rādītājiem un attēlošanas iestatījumi.
# Atslēga = rinda lapā "SECAP2030 mērķi".
TARGETS = {
  7:  dict(table="energy", rsec="pas", num=1,  code="1001",    unit="objekti",        dec=0),
  8:  dict(table="energy", rsec="pas", num=2,  code="1002",    unit="MWh",            dec=0),
  9:  dict(table="energy", rsec="pas", num=3,  code="3002",    unit="%",              dec=2),
  10: dict(table="energy", rsec="pas", num=4,  code="LV-RESE", unit="%",              dec=1),
  11: dict(table="energy", rsec="pas", num=5,  code="1012",    unit="MWh/gaismeklis", dec=3),
  12: dict(table="energy", rsec="pas", num=6,  code=None,      unit="",               dec=0),
  13: dict(table="energy", rsec="pas", num=7,  code="LV-RESE", unit="%",              dec=1),
  14: dict(table="energy", rsec="pas", num=8,  code="1015",    unit="%",              dec=1),
  16: dict(table="energy", rsec="dz",  num=10, code="2004",    unit="ēkas",           dec=0),
  17: dict(table="energy", rsec="dz",  num=11, code="2001",    unit="MWh",            dec=0,
          div=1000, dunit="GWh", bval="2456 GWh (2123 GWh)"),
  19: dict(table="energy", rsec="en",  num=12, code="3002",    unit="%",              dec=2),
  20: dict(table="energy", rsec="en",  num=13, code=None,      unit="%",              dec=2),
  22: dict(table="energy", rsec="en",  num=14, code="3006",    unit="%",              dec=1),
  23: dict(table="energy", rsec="en",  num=15, code="3007",    unit="%",              dec=1),
  24: dict(table="energy", rsec="en",  num=16, code="3005",    unit="GWh",            dec=0,
          bval="3461 GWh (362 429 tūkst. m3)"),
  26: dict(table="energy", rsec="en",  num=17, code="3009",    unit="%",              dec=1),
  28: dict(table="energy", rsec="tr",  num=19, code="4007",    unit="%",              dec=1),
  29: dict(table="energy", rsec="tr",  num=20, code=None,      unit="",               dec=0),
  34: dict(table="co2",    rsec="kop", num=0,  code="5003",    unit="%",              dec=1),
  35: dict(table="co2",    rsec="en",  num=18, code="5005",    unit="%",              dec=1),
  36: dict(table="co2",    rsec="tr",  num=21, code="5007",    unit="%",              dec=1),
  37: dict(table="co2",    rsec="pas", num=9,  code="5008",    unit="%",              dec=1),
}
# Rādītāju papildu iestatījumi (decimāldaļas, piezīmes, datu rindas pārrāvums)
IND_EXTRA = {
  "3002": dict(dec=2),
  "4005": dict(dec=1),
  "2004": dict(breakBefore=2025,
               notes={2025: "Ekonomikas ministrijas dati; nav salīdzināms ar 2020–2024 vērtībām (cita uzskaite)"}),
  "LV-RESE": dict(notes={2025: "provizoriski"}),
}
SECTOR_OF_IND = {"1": "pas", "2": "dz", "3": "en", "4": "tr", "5": "kop"}
SECTOR_NAMES = {"pas": "Pašvaldības infrastruktūra", "dz": "Daudzdzīvokļu ēkas",
                "en": "Enerģijas ražošana", "tr": "Transports", "kop": "Rīgas kopējās emisijas"}
SITE = dict(
  title="Rīgas SECAP vidusposma pārskats 2026",
  version_label="Darba versija",
  contact_email="krista.petersone@riga.lv",   # e-pasta adrese papildinājumiem
  repo_url="",        # GitHub repozitorija adrese (neobligāti)
  report_pdf="dati/Riga-SECAP-vidusposma-parskats-2026-darba-versija.pdf",
  license="CC BY 4.0",
  license_url="https://creativecommons.org/licenses/by/4.0/deed.lv",
  changelog=[["2026-10-21", "Pirmā darba versija kolēģu pārskatīšanai un papildināšanai."]],
)

# ---------------------------------------------------------------- palīgfunkcijas
def r4(x):
    return None if x is None else round(float(x), 4)

def num(x):
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None

def paras(s):
    return [p.strip() for p in re.split(r"\n\s*\n", str(s or "")) if p.strip()]

def clean(s):
    return re.sub(r"[ \t]+", " ", str(s or "")).strip()

# ---------------------------------------------------------------- rādītāji
def read_indicators(wb):
    ws = wb["Rādītāji - vērtības"]
    hdr = {ws.cell(6, c).value: c for c in range(1, ws.max_column + 1)}
    ycol = {y: hdr[y] for y in YEARS}
    meth = {}
    wm = wb["Rādītāji - apraksts metode"]
    for r in wm.iter_rows(min_row=3):
        code = r[1].value
        if code is None:
            continue
        meth[str(int(code)) if isinstance(code, (int, float)) else str(code)] = dict(
            desc=clean(r[7].value), link=clean(r[8].value), units=clean(r[10].value))
    out, group = [], None
    for r in range(7, ws.max_row + 1):
        code, name = ws.cell(r, 2).value, ws.cell(r, 3).value
        if name is None:
            continue
        if code is None and ws.cell(r, 4).value is None:
            group = clean(name); continue
        aux = code is None
        code = "LV-RESE" if aux else str(int(code))
        unit = clean(ws.cell(r, 6).value)
        vals = [ws.cell(r, ycol[y]).value for y in YEARS]
        vals = [None if (v in ("", None) or isinstance(v, str)) else float(v) for v in vals]
        pct = "%" in unit
        if pct:
            vals = [None if v is None else v * 100 for v in vals]
        if code.startswith("50") and code != "5001":
            unit = "% pret bāzi"
        ind = dict(code=code, name=clean(name), sector=clean(ws.cell(r, 4).value),
                   sub=clean(ws.cell(r, 5).value), unit=unit, group=group,
                   values=[r4(v) for v in vals], notes=[None] * 6, aux=aux,
                   rsec="en" if aux else SECTOR_OF_IND.get(code[0]))
        base = ws.cell(r, 7).value
        if isinstance(base, (int, float)) and code.startswith("50"):
            ind["baseKt"] = base
        m = meth.get(code)
        if m:
            ind["desc"], ind["link"], ind["units"] = m["desc"], m["link"], m["units"]
        ex = IND_EXTRA.get(code, {})
        for k, v in ex.items():
            if k == "notes":
                for y, t in v.items():
                    ind["notes"][YEARS.index(y)] = t
            else:
                ind[k] = v
        out.append(ind)
    return out

# ---------------------------------------------------------------- mērķi
def read_targets(wb):
    ws = wb["SECAP2030 mērķi"]
    out, sector = [], None
    for r in range(6, 38):
        a = ws.cell(r, 1).value
        if a and ws.cell(r, 6).value is None:
            sector = clean(a)
        if r not in TARGETS:
            continue
        cfg = TARGETS[r]
        g = lambda c: ws.cell(r, c).value
        byears = str(g(5) or "").split("\n")
        byc = int(re.sub(r"\D", "", byears[0]) or 2019)
        co2 = cfg["table"] == "co2"
        scale = 100 if (cfg["unit"] == "%" and cfg["code"] not in (None,)) or co2 else 1
        v2025, base, target = num(g(7)), num(g(10)), num(g(11))
        if cfg["code"] is None:   # mērķi bez rādītāja datu rindas
            scale = 100 if cfg["unit"] == "%" else 1
        t = dict(id=r, table=cfg["table"], sector="CO2 emisijas" if co2 else sector,
                 rsec=cfg["rsec"], num=cfg["num"], goal=clean(g(1)),
                 tval=" / ".join(clean(x) for x in str(g(2) or "").split("\n")),
                 tyear=clean(g(3)), bval=" / ".join(clean(x) for x in str(g(4) or "").split("\n")),
                 byear=" / ".join(clean(x) for x in str(g(5) or "").split("\n")),
                 status=clean(g(6)),
                 v2025=r4(v2025 * scale) if v2025 is not None else None,
                 progress=num(g(8)),
                 base=r4(base * scale) if base is not None else None,
                 target=r4(target * scale) if target is not None else None,
                 tp=num(g(12)) or (0.5 if byc == 2020 else 6 / 11), byc=byc,
                 code=cfg["code"], unit=cfg["unit"], dec=cfg["dec"], comment=clean(g(9)))
        if cfg["code"] is None and t["status"] != "Netiek vērtēts":
            t["base"] = t["target"] = None
        for k in ("div", "dunit", "bval"):
            if k in cfg:
                t[k] = cfg[k]
        out.append(t)
    return out

# ---------------------------------------------------------------- saturs
def sheet_rows(wb, name):
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in row):
            yield dict(zip(hdr, row))

def read_content(wb):
    statuses = [dict(name=clean(r["Statuss"]), desc=clean(r["Apraksts"])) for r in sheet_rows(wb, "Statusi")]
    measures = []
    for r in sheet_rows(wb, "Pasakumi"):
        projects = []
        for line in str(r["ES_projekti"] or "").split("\n"):
            if not line.strip():
                continue
            name, _, url = line.partition("|")
            projects.append(dict(name=clean(name), url=clean(url)))
        measures.append(dict(
            id=clean(r["ID"]), sector=clean(r["Sektors"]), title=clean(r["Nosaukums"]),
            status=clean(r["Statuss"]), status_note=clean(r["Statusa_pamatojums"]),
            check=bool(r["Parbaudit"]), plan=paras(r["Plans"]), done=paras(r["Paveiktais"]),
            indicators=[x.strip() for x in str(r["Saistitie_raditaji"] or "").split(",") if x.strip()],
            projects=projects, mwh=num(r["Ietaupijums_MWh"]), aer=num(r["AER_MWh"]),
            co2=num(r["CO2_samazinajums_t"]), inv=num(r["Investicijas_tukst_EUR"]),
            term=clean(r["Termins"]), shared=clean(r["Kopigs_ar"])))
    sectors = {}
    for r in sheet_rows(wb, "Sektori"):
        sectors[clean(r["ID"])] = dict(id=clean(r["ID"]), name=clean(r["Nosaukums"]),
            summary=paras(r["Kopsavilkums"]), intro=paras(r["Ievads"]),
            targets_text=paras(r["Merku_komentars"]), ind_text=paras(r["Raditaju_komentars"]),
            meas_intro=paras(r["Pasakumu_ievads"]), challenges=[], next=[])
    for r in sheet_rows(wb, "Izaicinajumi"):
        sectors[clean(r["Sektors"])]["challenges"].append(dict(
            group=clean(r["Grupa"]), y2022=clean(r["Izaicinajums_2022"]), y2026=clean(r["Komentars_2026"])))
    for r in sheet_rows(wb, "Nakamie_soli"):
        sectors[clean(r["Sektors"])]["next"].append(clean(r["Teksts"]))
    texts = {clean(r["Atslega"]): paras(r["Teksts"]) for r in sheet_rows(wb, "Teksti")}
    abbr = [[clean(r["Saisinajums"]), clean(r["Skaidrojums"])] for r in sheet_rows(wb, "Saisinajumi")]
    return dict(statuses=statuses, measures=measures, sectors=sectors, texts=texts, abbr=abbr)

# ---------------------------------------------------------------- lejupielādes
def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)

def downloads(data, outdir):
    outdir.mkdir(exist_ok=True)
    ind_h = ["kods", "nosaukums", "sektors", "apakšsektors", "vienība"] + [str(y) for y in YEARS] + \
            ["piezīmes", "apraksts", "avots"]
    ind_rows = []
    for i in data["indicators"]:
        notes = "; ".join(f"{YEARS[k]}: {n}" for k, n in enumerate(i["notes"]) if n)
        ind_rows.append([i["code"], i["name"], i["sector"], i["sub"], i["unit"]] + i["values"] +
                        [notes, i.get("desc", ""), i.get("link", "")])
    write_csv(outdir / "riga-secap-raditaji-2020-2025.csv", ind_h, ind_rows)
    tg_h = ["nr", "sektors", "mērķis", "mērķa vērtība", "mērķa gads", "bāzes vērtība", "bāzes gads",
            "2025. gada vērtība", "vienība", "progress", "laika progress", "statuss", "rādītājs", "komentārs"]
    tg_rows = [[t["num"] or "", SECTOR_NAMES[t["rsec"]], t["goal"], t["tval"], t["tyear"], t["bval"], t["byear"],
                t["v2025"], "% pret 2019" if t["table"] == "co2" else t["unit"],
                None if t["progress"] is None else round(t["progress"], 4), round(t["tp"], 4),
                t["status"], t["code"] or "", t["comment"]] for t in data["targets"]]
    write_csv(outdir / "riga-secap-merki-2025.csv", tg_h, tg_rows)
    ms_h = ["kods", "sektors", "pasākums", "statuss (darba versija)", "statusa pamatojums",
            "plānotais ietaupījums, MWh/gadā", "plānotā AER, MWh/gadā", "plānotais CO2 samazinājums, t/gadā",
            "indikatīvās investīcijas, tūkst. EUR", "ieviešanas termiņš", "kopīgs plāna rādītājs ar",
            "saistītie rādītāji", "ES sadarbības projekti", "plāns", "paveiktais"]
    ms_rows = [[m["id"], SECTOR_NAMES[m["sector"]], m["title"], m["status"], m["status_note"],
                m["mwh"], m["aer"], m["co2"], m["inv"], m["term"], m["shared"], ", ".join(m["indicators"]),
                "; ".join(p["name"] + (f" ({p['url']})" if p["url"] else "") for p in m["projects"]),
                "\n\n".join(m["plan"]), "\n\n".join(m["done"])] for m in data["measures"]]
    write_csv(outdir / "riga-secap-pasakumi.csv", ms_h, ms_rows)
    # viena Excel datne ar visām tabulām
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    info = wb.create_sheet("Par datiem")
    for line in [
        "Rīgas SECAP 2022–2030 vidusposma pārskats – atvērtie dati (darba versija)",
        f"Sagatavots: {data['meta']['built']}",
        f"Licence: {SITE['license']} ({SITE['license_url']})",
        "Avots: Rīgas valstspilsētas pašvaldības aģentūra “Rīgas enerģētikas aģentūra” (REA).",
        "Lapas: Rādītāji (2020–2025), Mērķi (2025. gada novērtējums), Pasākumi (SECAP2030 pasākumi un statuss).",
        "Procentu rādītāji norādīti procentos (piem. 46,85 = 46,85 %).",
        "Pasākumu statuss ir darba versija. Plānotā ietekme – no SECAP2030 pasākumu tabulas.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    info["A1"].font = Font(bold=True, size=13)
    for name, h, rows in [("Rādītāji", ind_h, ind_rows), ("Mērķi", tg_h, tg_rows), ("Pasākumi", ms_h, ms_rows)]:
        ws = wb.create_sheet(name)
        ws.append(h)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="000B40")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "C2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 14
        ws.column_dimensions["B"].width = 50 if name == "Rādītāji" else 24
        ws.column_dimensions["C"].width = 50
    wb.save(outdir / "Riga-SECAP-vidusposms-dati.xlsx")

# ---------------------------------------------------------------- galvenā daļa
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raditaji", default=str(ROOT / "avoti/Riga-SECAP-2025-Raditaji.xlsx"))
    ap.add_argument("--saturs", default=str(ROOT / "avoti/Riga-SECAP-saturs.xlsx"))
    a = ap.parse_args()
    wbi = openpyxl.load_workbook(a.raditaji, data_only=True)
    wbc = openpyxl.load_workbook(a.saturs, data_only=True)
    data = dict(years=YEARS, indicators=read_indicators(wbi), targets=read_targets(wbi))
    data.update(read_content(wbc))
    data["site"] = SITE
    data["sector_names"] = SECTOR_NAMES
    data["meta"] = dict(built=datetime.date.today().isoformat(), data_version="2026-10-21")
    # pārbaudes
    codes = {i["code"] for i in data["indicators"]}
    for m in data["measures"]:
        miss = [c for c in m["indicators"] if c not in codes]
        if miss:
            print(f"BRĪDINĀJUMS: {m['id']} atsaucas uz nezināmiem rādītājiem {miss}")
    allowed = {s["name"] for s in data["statuses"]}
    for m in data["measures"]:
        if m["status"] not in allowed:
            print(f"BRĪDINĀJUMS: {m['id']} statuss “{m['status']}” nav lapā Statusi")
    downloads(data, ROOT / "dati")
    files = {}
    for d in ("dati", "avoti"):
        for f in sorted((ROOT / d).glob("*")):
            if f.is_file():
                files[f"{d}/{f.name}"] = f.stat().st_size
    data["meta"]["files"] = files
    js = "/* Ģenerēts ar build/build.py – nelabot ar roku */\nwindow.SECAP=" + \
         json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    (ROOT / "assets").mkdir(exist_ok=True)
    (ROOT / "assets/data.js").write_text(js, encoding="utf-8")
    print(f"Gatavs: {len(data['indicators'])} rādītāji, {len(data['targets'])} mērķi, "
          f"{len(data['measures'])} pasākumi → assets/data.js, dati/")

if __name__ == "__main__":
    main()
